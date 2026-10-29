# export_results_workbook.py
#
# Writes every results table of Approach 1 into one Excel workbook, one sheet per table, so the
# numbers behind the complete report can be sorted, filtered and reused in slides. It reads the
# same result files as build_approach1_report.py (the per-dataset sweeps and analyses, the pooled
# mixture and noise sweeps, sequential training, the exported-head manifests, and the correction
# experiments) and writes results/approach1_results.xlsx next to the report. The first sheet
# lists every other sheet and what it holds.
# Usage: python -m experiments.approach1_latent_self_correction.export_results_workbook

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from experiments.approach1_latent_self_correction import build_approach1_report as rep
from experiments.approach1_latent_self_correction import build_corruption_report as corr
from experiments.approach1_latent_self_correction.build_mixture_report import NAMES, Results
from experiments.approach1_latent_self_correction.build_report import ORIGINAL_TEST

HERE = Path(__file__).parent
OUTPUT = HERE / "results" / "approach1_results.xlsx"
KINDS = ("pre", "verifier", "unified")
MODES = {"pre": ("main",), "verifier": ("main",), "unified": ("pre_generation", "joint")}
METRICS = ("correctness_auroc", "error_auprc", "brier", "ece_10_bin", "pearson_correctness")
SCORERS = ("token_confidence", "separate_pre_generation", "separate_verifier", "unified_pre_generation", "unified_joint")
SCORER_NAME = {"token_confidence": "Answer-token probability", "separate_pre_generation": "Separate pre-generation head", "separate_verifier": "Separate verifier head",
               "unified_pre_generation": "Unified, pre-generation mode", "unified_joint": "Unified, joint mode"}
DATASET_NAME = {"chartqa_full": "ChartQA", "textvqa": "TextVQA", "docvqa": "DocVQA", "scienceqa_img": "ScienceQA"}
GROUP_ORDER = ("old", "new", "old_linear", "new_linear", "new_full", "cell_plain", "cell_contrast", "cell_ground", "cell_full")


class Book:
    """A workbook that collects sheets and an index of what each one holds."""

    def __init__(self):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self.index = []

    def add(self, name: str, description: str, headers: list, rows: list) -> None:
        """Add one sheet with a bold header row, frozen panes, and sized columns."""
        sheet = self.workbook.create_sheet(name[:31])
        sheet.append(headers)
        for row in rows:
            sheet.append(["" if value is None else value for value in row])
        fill = PatternFill("solid", fgColor="DCE6F1")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(row[index - 1])) for row in rows[:200] if index - 1 < len(row)])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
        self.index.append([name[:31], description, len(rows)])

    def save(self, path: Path) -> None:
        """Write the index as the first sheet and save the file."""
        sheet = self.workbook.create_sheet("Index", 0)
        sheet.append(["Sheet", "What it holds", "Rows"])
        for row in self.index:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCE6F1")
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 110
        sheet.freeze_panes = "A2"
        path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(path)


def stat(node, field: str = "mean"):
    """A value from a {mean, std} node, or the node itself if it is a plain number."""
    if node is None:
        return None
    return node[field] if isinstance(node, dict) else node


def add_foundations(book: Book, r: Results, pd: dict) -> None:
    """Sheets for the datasets, the answer-token probability baseline, and the original ChartQA study."""
    rows = []
    for key, name in rep.DATASETS:
        s = pd[key]["analysis"]["splits"]
        rows.append([name, s["train"]["count"], s["val"]["count"], s["test"]["count"], s["train"]["accuracy"], s["val"]["accuracy"], s["test"]["accuracy"], rep.RULE[key]])
    comp = r.analysis["composition"]
    rows.append(["Pooled mixture", comp["train"]["rows"], comp["val"]["rows"], comp["test"]["rows"], comp["train"]["accuracy"], comp["val"]["accuracy"], comp["test"]["accuracy"], "union of the four"])
    book.add("Datasets", "Split sizes and base accuracy of the Qwen3-VL-2B answers, per dataset and pooled", ["Dataset", "Train rows", "Validation rows", "Test rows", "Base accuracy, train", "Base accuracy, validation", "Base accuracy, test", "Correct means"], rows)

    rows = []
    for key, name in rep.DATASETS:
        a = pd[key]["analysis"]
        t = a["token_confidence_metrics"]
        rows.append([name, a["splits"]["test"]["accuracy"], a["token_confidence_exactly_one"], t["correctness_auroc"], t["error_auprc"], t["brier"], t["ece_10_bin"], t["pearson_correctness"], a["answer_length_baseline"]["auroc"]])
    p = r.baseline("pooled")
    rows.append(["Pooled mixture", comp["test"]["accuracy"], p["exactly_one"], p["token_confidence"]["correctness_auroc"], p["token_confidence"]["error_auprc"], p["token_confidence"]["brier"], p["token_confidence"]["ece_10_bin"], p["token_confidence"]["pearson_correctness"], p["answer_length"]["auroc"]])
    book.add("Token baseline", "The untrained answer-token probability and an answer-length baseline, test split", ["Dataset", "Base accuracy", "Share exactly 1.0", "AUROC", "Error AUPRC", "Brier", "ECE", "Pearson r", "Answer-length AUROC"], rows)

    book.add("Original study test", "Original ChartQA study, test split, 2,500 questions (from the original report)", ["Metric", "Answer-token probability", "Pre-gen MLP", "Pre-gen linear probe", "Verifier MLP (swept)", "Verifier, early stop (3 seeds)"], [[m] + v for m, v in ORIGINAL_TEST.items()])
    runs = rep.load_json(rep.SWEEP / "chartqa_original_validation_runs.json")
    book.add("Original sweep runs", "All 72 validation runs of the original ChartQA sweep", ["Head", "Architecture", "Learning rate", "Ranking weight", "Seed", "Error AUPRC", "AUROC", "Brier", "ECE", "Pearson r", "Parameters"],
             [[x["head"], x["architecture"], x["learning_rate"], x["ranking_weight"], x["seed"], x["error_auprc"], x["auroc"], x["brier"], x["ece"], x["pearson"], rep.params("pre" if x["head"] == "pre_generation" else "verifier", x["architecture"])] for x in runs])


def add_heads(book: Book, pd: dict) -> None:
    """Sheets for head sizes and the winning configuration of each per-dataset sweep."""
    rows = []
    for kind in KINDS:
        for arch in ("linear", "mlp_256", "mlp_512"):
            n = rep.params(kind, arch)
            rows.append([rep.KIND_NAME[kind], arch, rep.FEATURE if kind != "unified" else 2 * rep.FEATURE, n, n * 4 / 1e6, 100 * n / 2.0e9])
    book.add("Head sizes", "Exact parameter counts of every head architecture (hidden size 2048)", ["Head", "Architecture", "Input size", "Parameters", "Megabytes (fp32)", "Percent of a 2B-parameter model"], rows)
    rows = []
    for key, name in rep.DATASETS:
        s, u = pd[key]["separate"], pd[key]["unified"]
        pc, vc, uc = s["pre_generation"]["best_config"], s["candidate_verifier"]["best_config"], u["winner"]["config"]
        rows.append([name, "Separate pre-generation", pc["architecture"], pc["learning_rate"], pc["ranking_weight"], "", "", pc["seed"], rep.params("pre", pc["architecture"])])
        rows.append([name, "Separate verifier", vc["architecture"], vc["learning_rate"], vc["ranking_weight"], "", "", vc["seed"], rep.params("verifier", vc["architecture"])])
        rows.append([name, "Unified", uc["architecture"], uc["learning_rate"], uc["rank_weight"], uc["pre_weight"], uc["pair_weight"], "mean of 3 seeds", rep.params("unified", uc["architecture"])])
    book.add("Winners per dataset", "Winning configuration of the per-dataset sweeps (four-dataset study)", ["Dataset", "Head", "Architecture", "Learning rate", "Ranking weight", "Pre-generation weight", "Pair weight", "Seed", "Parameters"], rows)


def add_per_dataset_results(book: Book, pd: dict) -> None:
    """Sheets for per-dataset metrics, the gate curves, bootstrap intervals, and reliability bins."""
    rows = []
    for key, name in rep.DATASETS:
        for scorer in SCORERS:
            m = pd[key]["metrics"][scorer]
            rows.append([name, SCORER_NAME[scorer]] + [m.get(metric) for metric in METRICS])
    book.add("Per-dataset metrics", "Test metrics of the five scorers on each dataset (unified rows are means over three seeds)", ["Dataset", "Scorer", "AUROC", "Error AUPRC", "Brier", "ECE", "Pearson r"], rows)
    rows = []
    for key, name in rep.DATASETS:
        a = pd[key]["analysis"]
        for scorer in SCORERS:
            for i, fraction in enumerate(a["grid"]):
                rows.append([name, SCORER_NAME[scorer], fraction, a["scorers"][scorer]["accuracy_at_coverage"][i], a["scorers"][scorer]["error_capture"][i]])
    book.add("Gate curves", "Accuracy of the most trusted fraction kept, and share of all errors captured when the same fraction of least trusted answers is flagged", ["Dataset", "Scorer", "Fraction", "Accuracy of kept fraction (coverage)", "Errors captured (flag rate)"], rows)
    rows = []
    for key, name in rep.DATASETS:
        for comparison, label in (("joint_minus_separate_verifier", "Unified joint minus separate verifier"), ("unified_pre_minus_separate_pre", "Unified pre-generation minus separate pre-generation")):
            for metric in ("auroc", "error_auprc"):
                b = pd[key]["analysis"]["bootstrap"][comparison][metric]
                rows.append([name, label, metric, b["difference"], b["ci_low"], b["ci_high"]])
    book.add("Bootstrap", "Paired bootstrap differences (1,000 resamples) with 95% intervals", ["Dataset", "Comparison", "Metric", "Difference", "CI low", "CI high"], rows)
    rows = []
    for key, name in rep.DATASETS:
        for scorer in ("token_confidence", "unified_joint"):
            for b in pd[key]["analysis"]["reliability"][scorer]:
                rows.append([name, SCORER_NAME[scorer], b["low"], b["count"], b["mean_probability"], b["accuracy"]])
    book.add("Reliability bins", "Calibration bins: mean predicted probability against observed accuracy", ["Dataset", "Scorer", "Bin lower edge", "Count", "Mean predicted probability", "Observed accuracy"], rows)
    rows = []
    for key, name in rep.DATASETS:
        for group in ("caught_errors", "missed_errors", "false_alarms", "confident_correct"):
            for case in pd[key]["analysis"]["case_studies"][group]:
                rows.append([name, group, case["index"], case["question"][:300], case["reference"], str(case["prediction"])[:200], case["joint_score"], case["pre_score"], case["token_confidence"]])
    book.add("Case studies", "Examples the unified head caught, missed, falsely flagged, or trusted correctly", ["Dataset", "Group", "Index", "Question", "Reference", "Model answer", "Unified joint score", "Pre-generation score", "Answer-token probability"], rows)


def add_mixture(book: Book, r: Results) -> None:
    """Sheets for the pooled mixture under the old and new loss, swap and noise, with grounding."""
    for label, loader, sheet_prefix in (("swap", lambda k: corr.load(f"mixture_{k}"), "Mixture swap"), ("noise", lambda k: corr.load(f"mixture_noise_{k}"), "Mixture noise")):
        rows, grounding_rows, config_rows, sweeps = [], [], [], {}
        for kind in KINDS:
            sweep = loader(kind)
            if sweep is None:
                continue
            sweeps[kind] = sweep
            for group in GROUP_ORDER:
                w = sweep["winners"].get(group)
                if w is None:
                    continue
                c = w["config"]
                config_rows.append([kind, group, c["architecture"], c["learning_rate"], c.get("rank_weight", 0), c.get("pre_weight", 0), c.get("pair_weight", 0), c["ground_weight"], c["ground_on"], rep.params(kind, c["architecture"]), json.dumps(w["epochs"])])
                for mode in MODES[kind]:
                    for dataset in ["pooled"] + r.datasets:
                        node = w["test"][mode] if dataset == "pooled" else w["test_per_dataset"].get(dataset, {}).get(mode)
                        if node is None:
                            continue
                        row = [kind, group, mode, NAMES.get(dataset, "Pooled")]
                        for metric in METRICS:
                            row += [stat(node.get(metric)), stat(node.get(metric), "std")]
                        rows.append(row)
                    for dataset in ["pooled"] + r.datasets:
                        g = w["grounding"][mode].get(dataset)
                        if g:
                            for subset in ("correct", "all"):
                                grounding_rows.append([kind, group, mode, label + " (training corruption)", NAMES.get(dataset, "Pooled"), subset, stat(g[subset]["order_accuracy"]), stat(g[subset]["order_accuracy"], "std"), stat(g[subset]["mean_drop"]), stat(g[subset]["count"])])
                    for corruption, by_mode in (w.get("grounding_extra") or {}).items():
                        for dataset, g in by_mode[mode].items():
                            grounding_rows.append([kind, group, mode, corruption + " (graded only)", NAMES.get(dataset, "Pooled") if dataset != "pooled" else "Pooled", "correct", stat(g["correct"]["order_accuracy"]), stat(g["correct"]["order_accuracy"], "std"), stat(g["correct"]["mean_drop"]), stat(g["correct"]["count"])])
        headers = ["Head kind", "Winner group", "Mode", "Dataset"]
        for metric in METRICS:
            headers += [metric + " mean", metric + " std"]
        book.add(f"{sheet_prefix} metrics", f"Pooled-mixture test metrics of every winner group, {label} training corruption, mean and std over three seeds", headers, rows)
        book.add(f"{sheet_prefix} grounding", f"Share of answers whose score falls under corruption, {label} version, pooled and per dataset, with other corruptions where graded", ["Head kind", "Winner group", "Mode", "Corruption", "Dataset", "Subset", "Order accuracy mean", "Order accuracy std", "Mean drop", "Count"], grounding_rows)
        book.add(f"{sheet_prefix} winners", f"Winning configuration of every group, {label} version", ["Head kind", "Winner group", "Architecture", "Learning rate", "Rank weight", "Pre-generation weight", "Pair weight", "Grounding weight", "Grounded on", "Parameters", "Best epoch per seed"], config_rows)
        rows = []
        for kind, sweep in sweeps.items():
            for item in sweep["configs"]:
                c = item["config"]
                for mode, metrics in item["val"].items():
                    rows.append([kind, c["architecture"], c["learning_rate"], c.get("rank_weight", 0), c.get("pre_weight", 0), c.get("pair_weight", 0), c["ground_weight"], mode] + [metrics[m] for m in METRICS])
        book.add(f"{sheet_prefix} sweep", f"Validation metrics of every configuration swept, {label} version (mean over three seeds)", ["Head kind", "Architecture", "Learning rate", "Rank weight", "Pre-generation weight", "Pair weight", "Grounding weight", "Mode", "AUROC", "Error AUPRC", "Brier", "ECE", "Pearson r"], rows)
    rows = []
    for kind in ("verifier", "unified"):
        mode = "main" if kind == "verifier" else "joint"
        for cell, name in (("cell_plain", "Neither term"), ("cell_contrast", "Pair term only"), ("cell_ground", "Grounding term only"), ("cell_full", "Both terms")):
            sweep = corr.load(f"mixture_{kind}")
            if sweep and cell in sweep["winners"]:
                w = sweep["winners"][cell]
                base = sweep["winners"]["cell_plain"]
                rows.append([rep.KIND_NAME[kind], name, w["test"][mode]["correctness_auroc"]["mean"], w["test"][mode]["correctness_auroc"]["std"], w["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["mean"], w["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["std"],
                             w["test"][mode]["correctness_auroc"]["mean"] - base["test"][mode]["correctness_auroc"]["mean"], w["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["mean"] - base["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["mean"]])
    book.add("Ablation 2x2", "Pair term and grounding term switched on and off (swap, pooled mixture)", ["Head", "Terms on", "AUROC mean", "AUROC std", "Order accuracy mean", "Order accuracy std", "AUROC change vs plain", "Order accuracy change vs plain"], rows)
    rows = []
    for key, name in (("chartqa", "ChartQA"), ("textvqa", "TextVQA"), ("docvqa", "DocVQA"), ("scienceqa", "ScienceQA")):
        for corruption in ("swap", "noise"):
            for kind in KINDS:
                sweep = corr.load(f"{key}_{corruption}_{kind}")
                if sweep is None:
                    continue
                for group in ("old", "new"):
                    w = sweep["winners"][group]
                    for mode in MODES[kind]:
                        rows.append([name, corruption, kind, group, mode, w["test"][mode]["correctness_auroc"]["mean"], w["test"][mode]["correctness_auroc"]["std"], w["test"][mode]["error_auprc"]["mean"],
                                     w["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["mean"], w["grounding"][mode]["pooled"]["correct"]["order_accuracy"]["std"], w["config"]["architecture"], rep.params(kind, w["config"]["architecture"])])
    book.add("Single-dataset losses", "Each dataset alone: old and new loss, swap and noise, AUROC and the share of correct answers whose score falls under the training corruption", ["Dataset", "Corruption", "Head kind", "Loss", "Mode", "AUROC mean", "AUROC std", "Error AUPRC mean", "Order accuracy mean", "Order accuracy std", "Architecture", "Parameters"], rows)
    rows = []
    for kind in KINDS:
        for label, tag in (("swap", "mixture_xeval"), ("noise", "mixture_noise")):
            sweep = corr.load(f"{tag}_{kind}")
            for group, group_label in (("old", "Old loss"), ("new", f"New loss, trained with {label}")):
                if sweep is None:
                    continue
                rows.append([rep.KIND_NAME[kind], group_label] + [corr.grounding_against(sweep, group, kind, c) for c in corr.CORRUPTIONS])
    book.add("Corruption transfer", "Share of correct answers whose score falls under each corruption (swap, noise, mask, shuffle), pooled test split", ["Head", "Trained with"] + list(corr.CORRUPTIONS), rows)


def add_sequential(book: Book, r: Results) -> None:
    """Sheets for sequential training summaries, stage matrices, and isolated references."""
    summary, matrices, isolated = [], [], []
    for label, loader in (("swap", lambda k: r.continual.get(k)), ("noise", lambda k: corr.load(f"mixture_noise_continual_{k}"))):
        for kind in KINDS:
            data = loader(kind)
            if not data:
                continue
            for variant, node in data["variants"].items():
                for mode in MODES[kind]:
                    s = node["summary"][mode]
                    final = [e["grounding_matrix"][mode][-1] for e in node["orders"].values()]
                    ground = sum(sum(row) / len(row) for row in final) / len(final)
                    summary.append([label, kind, mode, variant, s["final_avg"]["mean"], s["final_avg"]["std"], s["forgetting"]["mean"], s["forgetting"]["std"], s["bwt"]["mean"], s["bwt"]["std"], s["retention"]["mean"], s["retention"]["std"], ground])
                    for order, entry in node["orders"].items():
                        for stage, row in enumerate(entry["matrix"][mode]):
                            for d, dataset in enumerate(data["datasets"]):
                                matrices.append([label, kind, mode, variant, order, stage + 1, NAMES.get(entry["order"][stage], entry["order"][stage]), NAMES.get(dataset, dataset), row[d], entry["grounding_matrix"][mode][stage][d]])
                    for dataset, v in node["isolated"][mode].items():
                        isolated.append([label, kind, mode, variant, NAMES.get(dataset, dataset), v["mean"], v["std"]])
    book.add("Sequential summary", "Final average AUROC, forgetting, backward transfer, retention, and grounding after the last stage, per head, loss and replay", ["Corruption", "Head kind", "Mode", "Variant", "Final average mean", "std", "Forgetting mean", "std", "Backward transfer mean", "std", "Retention mean", "std", "Grounding after last stage"], summary)
    book.add("Sequential matrices", "AUROC and grounding of every dataset after every stage, per order (means over three seeds)", ["Corruption", "Head kind", "Mode", "Variant", "Order", "Stage", "Trained on at this stage", "Scored on", "AUROC", "Order accuracy"], matrices)
    book.add("Sequential isolated", "Head trained on each dataset alone, scored on that dataset (reference for retention)", ["Corruption", "Head kind", "Mode", "Variant", "Dataset", "AUROC mean", "AUROC std"], isolated)


def add_trained_heads(book: Book) -> None:
    """Sheets for the manifests of the exported heads."""
    root = rep.REPO / "trained_heads" / "approach1"
    rows = []
    for manifest in sorted((root / "winners").glob("*/manifest.json")):
        for e in json.loads(manifest.read_text(encoding="utf-8")):
            rows.append([manifest.parent.name, e["kind"], "+".join(e["groups"]), e["config"]["architecture"], e["config"]["learning_rate"], e["config"]["ground_weight"], ", ".join(e["files"]), *e["test_auroc_per_seed"], e["test_auroc_mean"], e["sweep_reported_mean"], e["difference"]])
    book.add("Trained heads winners", "Every exported winner: files, per-seed test AUROC, mean, sweep mean, and the difference", ["Tag", "Head kind", "Winner groups", "Architecture", "Learning rate", "Grounding weight", "Files", "Seed 1 AUROC", "Seed 2 AUROC", "Seed 3 AUROC", "Mean AUROC", "Sweep mean AUROC", "Difference"], rows)
    rows = []
    for manifest in sorted((root / "sequential").glob("*/manifest_*.json")):
        for e in json.loads(manifest.read_text(encoding="utf-8")):
            rows.append([manifest.parent.name, e["kind"], e["variant"], " > ".join(NAMES.get(n, n) for n in e["order"]), e["file"]] + [e["final_test_auroc"].get(d) for d, _ in rep.DATASETS])
    book.add("Trained heads sequential", "Final head of each sequential variant and order (seed 1) and its final test AUROC per dataset", ["Tag", "Head kind", "Variant", "Order", "File"] + [n for _, n in rep.DATASETS], rows)


def add_correction(book: Book) -> None:
    """Sheets for the correction experiments."""
    data = rep.correction_tables()
    if not data:
        return
    for key, sheet, text in (("retry", "Correction retry", "Prompted retry policies"), ("rerank", "Correction rerank", "Scored retry, reranking, gated margins and the oracle")):
        if key not in data:
            continue
        rows = []
        for policy, p in data[key]["policies"].items():
            pooled = p["pooled"]
            rows.append([policy, pooled["baseline_accuracy"], pooled["final_accuracy"], pooled["retried"], pooled["coverage"], pooled["fixed_errors"], pooled["broken_answers"], pooled["net_corrections"]] + [p["per_dataset"][d]["final_accuracy"] for d, _ in rep.DATASETS])
        book.add(sheet, text + " (100 examples, 25 per dataset)", ["Policy", "Baseline accuracy", "Final accuracy", "Retried", "Coverage", "Fixed", "Broken", "Net", "ChartQA", "TextVQA", "DocVQA", "ScienceQA"], rows)
    rows = []
    for policy, p in data["best"]["policies"].items():
        pooled = p["pooled"]
        low, high = rep.wilson(pooled["correct"], pooled["examples"])
        rows.append([policy, pooled["correct"], pooled["examples"], pooled["accuracy"], low, high] + [p["per_dataset"][d]["accuracy"] for d, _ in rep.DATASETS])
    book.add("Correction best-of-five", "Selection methods over five sampled candidates, with Wilson 95% intervals", ["Method", "Correct", "Examples", "Accuracy", "95% low", "95% high", "ChartQA", "TextVQA", "DocVQA", "ScienceQA"], rows)
    rows = []
    for policy, p in data["deep"]["policies"].items():
        pooled = p["pooled"]
        rows.append([policy, pooled.get("final_accuracy", pooled.get("accuracy")), pooled.get("fixed_errors", 0), pooled.get("broken_answers", 0)] + [p["per_dataset"][d].get("final_accuracy", p["per_dataset"][d].get("accuracy")) for d, _ in rep.DATASETS])
    book.add("Correction deeper pass", "Replacing the head-selected candidate with the deeper-pass answer", ["Policy", "Accuracy", "Fixed", "Broken", "ChartQA", "TextVQA", "DocVQA", "ScienceQA"], rows)
    rows = []
    for rec in data.get("deep_records", []):
        answer = rec["deep_answer"]["answer"]
        tokens = len(rec["deep_answer"]["token_probabilities"])
        rows.append([rec["dataset"], rec["index"], str(rec["reference_answer"])[:100], rec["selected_candidate"]["answer"][:100], rec["selected_candidate"]["correct"], tokens, tokens >= 64, answer.startswith("A candidate answer was"), answer[:300], rec["deep_answer"]["correct"]])
    book.add("Deeper pass per example", "Every deeper-pass answer with its length, whether it hit the 64-token limit, and whether it echoed the prompt", ["Dataset", "Index", "Reference", "Selected candidate", "Selected correct", "Deep answer tokens", "Hit 64-token limit", "Echoed the prompt", "Deep answer (first 300 chars)", "Deep answer correct"], rows)
    rows = []
    for rec in data.get("records", []):
        chosen = rec["selected_indices"]["head_best"]
        for k, c in enumerate(rec["candidates"]):
            rows.append([rec["dataset"], rec["index"], str(rec["reference_answer"])[:100], k + 1, c["answer"][:100], c["head_confidence"], c["token_confidence"], c["correct"], k == chosen, rec["greedy"]["correct"]])
    book.add("Best-of-five candidates", "All 500 sampled candidates with head score, answer-token probability and correctness", ["Dataset", "Index", "Reference", "Candidate", "Answer", "Head score", "Answer-token probability", "Correct", "Selected by head", "Greedy answer correct"], rows)


def add_full_correction(book: Book) -> None:
    """The correction experiments on the full test split with paired statistics against each baseline."""
    from experiments.approach1_latent_self_correction import correction_stats as st
    full = rep.full_correction_data()
    if not full:
        return
    def stat_row(name, base, final, extra=()):
        row = st.policy_row(base, final)
        return [name, *extra, row["examples"], row["accuracy"], row["accuracy_low"], row["accuracy_high"], row["difference"], row["ci_low"], row["ci_high"], row["fixed"], row["broken"], row["p_value"]]
    headers = ["Policy", "Questions", "Accuracy", "Accuracy 95% low", "Accuracy 95% high", "Difference from baseline", "Difference 95% low", "Difference 95% high", "Fixed", "Broken", "Sign-test p"]
    rows = []
    greedy = st.best5_outcomes(full["best5"], "greedy")
    for policy in st.BEST5_POLICIES:
        rows.append(stat_row(policy, greedy, st.best5_outcomes(full["best5"], policy)))
    book.add("Full best-of-five", "Best-of-five selection on the full test split (9,517 questions) against greedy, paired statistics", headers, rows)
    rows = [stat_row(name, *st.retry_outcomes(full["retry"], key)) for name, key in st.RETRY_POLICIES.items()]
    book.add("Full retry", "Prompted retry, reranking and gated policies on the full test split against no retry", headers, rows)
    rows = [stat_row(name, *st.deep_outcomes(full["deep"], key)) for name, key in st.DEEP_POLICIES.items()]
    book.add("Full deeper pass", "Deeper pass after best-of-five on the full test split against the selected candidate", headers, rows)
    rows = []
    for dataset, label in rep.DATASETS:
        for experiment, data, outcomes in (("best-of-five", full["best5"], None), ("retry", full["retry"], None), ("deeper pass", full["deep"], None)):
            sub = [r for r in data if r["dataset"] == dataset]
            if experiment == "best-of-five":
                base = st.best5_outcomes(sub, "greedy")
                for policy in st.BEST5_POLICIES[1:]:
                    rows.append(stat_row(policy, base, st.best5_outcomes(sub, policy), (label, experiment)))
            elif experiment == "retry":
                for policy, key in st.RETRY_POLICIES.items():
                    if policy != "baseline_no_retry":
                        rows.append(stat_row(policy, *st.retry_outcomes(sub, key), (label, experiment)))
            else:
                for policy, key in st.DEEP_POLICIES.items():
                    if policy != "keep_best_of_five":
                        rows.append(stat_row(policy, *st.deep_outcomes(sub, key), (label, experiment)))
    book.add("Full per dataset", "Every full-split correction policy by dataset, against its baseline", ["Policy", "Dataset", "Experiment"] + headers[1:], [[r[0], r[1], r[2]] + r[3:] for r in rows])
    failures = st.format_failures(full["deep"])
    book.add("Full deeper pass formats", "Deeper-pass answers that hit the 64-token limit or repeated the prompt, and the correct answers they broke", ["Questions", "Hit the token limit", "Repeated the prompt (not capped)", "Broken, capped", "Broken, repeated prompt"],
             [[failures["examples"], failures["capped"], failures["echoed"], failures["broken_capped"], failures["broken_echoed"]]])


def add_timings(book: Book) -> None:
    """Measured wall-clock times from Slurm accounting."""
    rows = [
        ["Corrupted-image features, ChartQA train (28,299 questions)", "1 GPU", 82.5],
        ["Pooled sweep, pre-generation head (36 configs x 3 seeds)", "1 GPU", 3.4],
        ["Pooled sweep, verifier head (36 configs x 3 seeds)", "1 GPU", 8.0],
        ["Pooled sweep, unified head (72 configs x 3 seeds)", "1 GPU", 11.6],
        ["Sequential training, pre-generation head (4 variants, 3 orders, 3 seeds)", "1 GPU", 1.5],
        ["Sequential training, verifier head", "1 GPU", 3.4],
        ["Sequential training, unified head", "1 GPU", 3.4],
        ["Export of the 48 pooled-mixture winner heads (retraining)", "1 GPU", 4.3],
    ]
    book.add("Timings", "Wall-clock minutes from Slurm accounting on the ICE cluster (one GPU each)", ["Job", "Resources", "Minutes"], rows)


def add_heads_vs_probes(book: Book) -> None:
    """The trained-heads-against-linear-probes comparison with every metric."""
    from experiments.approach1_latent_self_correction import heads_vs_probes as hp
    rows = []
    for r in hp.comparison_rows():
        row = [r["setting"], r["dataset"], r["head_label"], r["token_auroc"], r["linear_arch"], r["linear_params"], r["head_arch"], r["head_params"]]
        for metric in hp.METRICS:
            row += [r[f"linear_{metric}"], r[f"linear_{metric}_std"], r[f"head_{metric}"], r[f"head_{metric}_std"], r[f"diff_{metric}"]]
        row.append(hp.verdict(r))
        rows.append(row)
    headers = ["Setting", "Dataset", "Head", "Answer-token probability AUROC", "Linear probe architecture", "Linear probe parameters", "Trained head architecture", "Trained head parameters"]
    for metric in hp.METRICS:
        headers += [f"Linear {metric}", "std", f"Head {metric}", "std", f"Difference {metric}"]
    headers.append("Reading")
    book.add("Heads vs linear probes", "Best linear probe against best trained head per dataset, head kind and mode, pooled and dataset-alone training, all metrics", headers, rows)


def build_workbook(path: Path = OUTPUT) -> Path:
    """Build the workbook and return its path."""
    r = Results()
    pd = rep.per_dataset()
    book = Book()
    add_foundations(book, r, pd)
    add_heads(book, pd)
    add_per_dataset_results(book, pd)
    add_heads_vs_probes(book)
    add_mixture(book, r)
    add_sequential(book, r)
    add_trained_heads(book)
    add_correction(book)
    add_full_correction(book)
    add_timings(book)
    book.save(path)
    return path


def main() -> None:
    """Write the workbook and report its size and sheet count."""
    path = build_workbook()
    from openpyxl import load_workbook
    sheets = load_workbook(path, read_only=True).sheetnames
    print(f"[xlsx] wrote {path} ({path.stat().st_size / 1e3:.0f} KB, {len(sheets)} sheets)")


if __name__ == "__main__":
    main()
