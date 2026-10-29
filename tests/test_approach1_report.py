# test_approach1_report.py
#
# Offline tests for the complete Approach 1 report builder: exact head parameter counts, the
# template filler, the interval and AUROC helpers, and that the vendored KaTeX files are present
# and get embedded. The full report build needs the real result files, so it is not run here.
#
# Run with `python -m tests.test_approach1_report`, or with pytest.

import math

from experiments.approach1_latent_self_correction import build_approach1_report as report


def test_parameter_counts_match_the_instantiated_heads():
    """The closed-form counts equal the sizes of the real heads (checked against PyTorch once)."""
    assert report.params("pre", "linear") == 2049
    assert report.params("pre", "mlp_256") == 528897
    assert report.params("verifier", "mlp_512") == 1053697
    assert report.params("unified", "linear") == 12290
    assert report.params("unified", "mlp_256") == 1057537
    assert report.params("unified", "mlp_512") == 2106881


def test_fill_replaces_markers_and_keeps_latex_braces():
    """Markers are substituted, LaTeX braces are untouched, and a missing value is an error."""
    text = report.fill(r"\frac{1}{@@n@@} for @@name@@", n=4, name="x")
    assert text == r"\frac{1}{4} for x"
    try:
        report.fill("@@missing@@")
    except KeyError:
        return
    raise AssertionError("expected KeyError")


def test_wilson_interval_and_auroc_helper():
    """The interval brackets the proportion, and AUROC handles perfect, reversed and tied scores."""
    low, high = report.wilson(81, 100)
    assert low < 0.81 < high and 0.70 < low and high < 0.90
    assert report.auroc_of([0, 0, 1, 1], [0.1, 0.2, 0.8, 0.9]) == 1.0
    assert report.auroc_of([1, 1, 0, 0], [0.1, 0.2, 0.8, 0.9]) == 0.0
    assert math.isclose(report.auroc_of([0, 1, 0, 1], [0.5, 0.5, 0.5, 0.5]), 0.5)


def test_katex_is_embedded_with_fonts():
    """The page carries KaTeX CSS with data-URI fonts and the render script, so it works offline."""
    page = report.katex_inline()
    assert "data:font/woff2;base64," in page and "renderMathInElement" in page and "https://" not in page.split("<script>")[0]


def test_workbook_book_writes_sheets_and_an_index():
    """The workbook helper writes one sheet per table, truncates long names, and puts an index first."""
    import tempfile
    from pathlib import Path
    from openpyxl import load_workbook
    from experiments.approach1_latent_self_correction.export_results_workbook import Book
    book = Book()
    book.add("A table with a very long name that exceeds the limit", "first", ["x", "y"], [[1, 2], [3, None]])
    book.add("Second", "second", ["a"], [["z"]])
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "out.xlsx"
        book.save(path)
        wb = load_workbook(path)
    assert wb.sheetnames[0] == "Index" and len(wb.sheetnames) == 3
    assert all(len(name) <= 31 for name in wb.sheetnames)
    assert wb["Second"]["A2"].value == "z" and wb[wb.sheetnames[1]]["B3"].value in (None, "")


def _run_all() -> None:
    """Run every test_* function in this module and report pass and fail counts."""
    tests = [obj for name, obj in globals().items() if name.startswith("test_") and callable(obj)]
    failures = 0
    for test in tests:
        try:
            test()
            print(f"PASS: {test.__name__}")
        except AssertionError as exc:
            failures += 1
            print(f"FAIL: {test.__name__}: {exc}")
    print(f"\n{len(tests) - failures}/{len(tests)} passed")
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    _run_all()
